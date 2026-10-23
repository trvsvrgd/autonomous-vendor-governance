from datetime import datetime, timezone
import io
import logging
from typing import Any, Optional
from uuid import UUID

from docx import Document as DocxDocument
from pypdf import PdfReader
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.aws_clients import archive_payload_to_s3, publish_sqs_message, read_payload_from_s3
from app.config import get_settings
from app.models import IngestionSource, SourceType, VendorSnapshot
from app.workers.safe_fetch import compute_normalized_sha256

logger = logging.getLogger(__name__)


def extract_document_text_and_tables(file_bytes: bytes, filename: str) -> str:
    """Extract full text and tabular disclosures from PDF or DOCX contract/ULA/SOC2 files."""
    lower_name = filename.lower()

    if lower_name.endswith(".pdf") or file_bytes.startswith(b"%PDF"):
        reader = PdfReader(io.BytesIO(file_bytes))
        pages_text: list[str] = []
        for idx, page in enumerate(reader.pages, start=1):
            page_text = (page.extract_text() or "").strip()
            if page_text:
                pages_text.append(f"[Page {idx}]\n{page_text}")
        extracted = "\n\n".join(pages_text).strip()
        if not extracted:
            raise ValueError(f"PDF document '{filename}' contained no extractable text layer.")
        return extracted

    if lower_name.endswith(".docx"):
        doc = DocxDocument(io.BytesIO(file_bytes))
        blocks: list[str] = []
        for para in doc.paragraphs:
            text = para.text.strip()
            if text:
                blocks.append(text)
        for t_idx, table in enumerate(doc.tables, start=1):
            table_rows: list[str] = [f"[Table {t_idx} Disclosure]"]
            for row in table.rows:
                cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
                if any(cells):
                    table_rows.append(" | ".join(cells))
            if len(table_rows) > 1:
                blocks.append("\n".join(table_rows))
        extracted = "\n\n".join(blocks).strip()
        if not extracted:
            raise ValueError(f"DOCX document '{filename}' contained no extractable text or tables.")
        return extracted

    # Plaintext / Markdown contract addendum fallback
    decoded = file_bytes.decode("utf-8", errors="replace").strip()
    if not decoded:
        raise ValueError(f"Uploaded document '{filename}' is empty.")
    return decoded


def ingest_contract_or_ula_document(
    db: Session,
    *,
    vendor_id: UUID,
    source_type: SourceType,
    filename: str,
    file_bytes: Optional[bytes] = None,
    s3_key_or_uri: Optional[str] = None,
    source_id: Optional[UUID] = None,
) -> dict[str, Any]:
    """Contract and ULA Document Ingestion Worker.

    Triggers when a PDF or DOCX file (enterprise ULA, AI addendum, or SOC 2 report) is uploaded
    to the encrypted S3 document bucket. Extracts full text and tabular disclosures, writes a
    new record to vendor_snapshots, and pushes an evaluation job to the SQS queue.
    """
    settings = get_settings()
    if file_bytes is None:
        if not s3_key_or_uri:
            raise ValueError("Either file_bytes or s3_key_or_uri must be provided.")
        file_bytes = read_payload_from_s3(s3_key_or_uri)

    extracted_text = extract_document_text_and_tables(file_bytes, filename)
    sha256_hash = compute_normalized_sha256(extracted_text)
    now = datetime.now(timezone.utc)

    s3_archive_uri = archive_payload_to_s3(
        vendor_id=vendor_id,
        category=source_type.value,
        filename=filename,
        content_bytes=file_bytes,
        content_type="application/pdf"
        if filename.lower().endswith(".pdf")
        else "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )

    source: Optional[IngestionSource] = None
    if source_id:
        source = db.scalar(
            select(IngestionSource).where(
                IngestionSource.source_id == source_id,
                IngestionSource.vendor_id == vendor_id,
            )
        )
    if source is None:
        source = db.scalar(
            select(IngestionSource).where(
                IngestionSource.vendor_id == vendor_id,
                IngestionSource.source_type == source_type,
            )
        )
    if source is None:
        source = IngestionSource(
            vendor_id=vendor_id,
            source_type=source_type,
            target_uri_or_s3_key=s3_archive_uri,
            polling_cron_schedule="s3:ObjectCreated:*",
            last_polled_at=now,
            last_sha256_hash=sha256_hash,
        )
        db.add(source)
        db.flush()

    previous_snapshot = db.scalar(
        select(VendorSnapshot)
        .where(VendorSnapshot.source_id == source.source_id)
        .order_by(VendorSnapshot.captured_at.desc())
        .limit(1)
    )

    source.target_uri_or_s3_key = s3_archive_uri
    source.last_polled_at = now
    source.last_sha256_hash = sha256_hash

    snapshot = VendorSnapshot(
        source_id=source.source_id,
        vendor_id=vendor_id,
        raw_extracted_text=extracted_text,
        content_sha256_hash=sha256_hash,
        s3_archive_uri=s3_archive_uri,
        captured_at=now,
    )
    db.add(snapshot)
    db.commit()
    db.refresh(snapshot)

    evaluation_job = {
        "event_type": "CONTRACT_DOCUMENT_INGESTED",
        "vendor_id": str(vendor_id),
        "source_id": str(source.source_id),
        "previous_snapshot_id": str(previous_snapshot.snapshot_id) if previous_snapshot else None,
        "new_snapshot_id": str(snapshot.snapshot_id),
        "content_sha256_hash": sha256_hash,
        "captured_at": now.isoformat(),
    }
    sqs_msg_id = publish_sqs_message(settings.SQS_EVALUATION_QUEUE_URL, evaluation_job)

    return {
        "status": "ingested",
        "vendor_id": str(vendor_id),
        "source_id": str(source.source_id),
        "previous_snapshot_id": str(previous_snapshot.snapshot_id) if previous_snapshot else None,
        "new_snapshot_id": str(snapshot.snapshot_id),
        "s3_archive_uri": s3_archive_uri,
        "content_sha256_hash": sha256_hash,
        "extracted_characters": len(extracted_text),
        "sqs_message_id": sqs_msg_id,
    }
